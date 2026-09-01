"""Export renderers: typed Excel cells and right-to-left Hebrew PDF (Requirement 19.1, 19.2, 19.3).

These are pure tests of `app.export` — no database, no request. They pin the two claims the requirement
turns on:

* **Excel cells are typed, not text** (Requirement 19.1). A money figure lands in the workbook as a
  number a formula can sum, and a date as a real date, not a string that looks like one. The test
  loads the rendered bytes back with openpyxl and asserts each cell's Python type and number format.
* **Hebrew renders right-to-left in the PDF** (Requirement 19.3). WeasyPrint needs native libraries
  absent from a bare developer machine, so the byte-level render is asserted only where they are
  present; but the RTL correctness lives in the HTML WeasyPrint paints — the `dir="rtl"`, the `lang`,
  the Hebrew text — and that HTML is pure Python, so it is asserted unconditionally.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest

from app.export import (
    CellKind,
    ExportColumn,
    ExportDocument,
    ExportStamp,
    ExportTable,
    PdfRenderUnavailable,
    cell,
    to_pdf,
    to_pdf_html,
    to_xlsx,
)

# A Hebrew employee name and header, so the tests exercise real Hebrew rather than a placeholder.
_HEB_NAME = "\u05d0\u05d1\u05e8\u05d4\u05dd"  # Abraham
_HEB_EMPLOYEE = "\u05e2\u05d5\u05d1\u05d3"  # Employee
_HEB_COST = "\u05e2\u05dc\u05d5\u05ea"  # Cost


def _sample_document(language: str = "he") -> ExportDocument:
    return ExportDocument(
        title="\u05d3\u05d5\u05f4\u05d7 \u05e9\u05e2\u05d5\u05ea",
        title_en="Hours report",
        stamp=ExportStamp(
            period_label="08/2025",
            generated_at=datetime(2025, 9, 1, 10, 30),
            generated_by="admin",
            filters=[("\u05d0\u05ea\u05e8", "Site A")],
            currency="ILS",
        ),
        tables=[
            ExportTable(
                columns=[
                    ExportColumn(_HEB_EMPLOYEE, "Employee"),
                    ExportColumn(_HEB_COST, "Cost", CellKind.MONEY),
                    ExportColumn("\u05e9\u05e2\u05d5\u05ea", "Hours", CellKind.HOURS),
                    ExportColumn("\u05ea\u05d0\u05e8\u05d9\u05da", "Date", CellKind.DATE),
                    ExportColumn("\u05d3\u05e7\u05d5\u05ea", "Minutes", CellKind.INTEGER),
                ],
                rows=[
                    [
                        cell(_HEB_NAME),
                        cell(Decimal("7420.00"), CellKind.MONEY),
                        cell(Decimal("212.00"), CellKind.HOURS),
                        cell(date(2025, 8, 4), CellKind.DATE),
                        cell(12720, CellKind.INTEGER),
                    ],
                ],
                totals=[
                    cell("\u05e1\u05d4\u05f4\u05db"),
                    cell(Decimal("7420.00"), CellKind.MONEY),
                    cell(None, CellKind.HOURS),
                    cell(None, CellKind.DATE),
                    cell(12720, CellKind.INTEGER),
                ],
            )
        ],
        language=language,
    )


# ===================================================================== Excel: typed cells (19.1)


def test_excel_money_cell_is_a_number_not_text():
    """Requirement 19.1: a money figure is a numeric cell a formula can sum, not a string."""
    workbook = openpyxl.load_workbook(io.BytesIO(to_xlsx(_sample_document())))
    sheet = workbook.active

    money_cells = [c for row in sheet.iter_rows() for c in row if c.value == 7420.0]
    assert money_cells, "the money value did not land in any cell"
    for money in money_cells:
        assert isinstance(money.value, int | float), "money must be a number, not text"
        assert "0.00" in money.number_format, "money keeps a two-place currency format"


def test_excel_date_cell_is_a_real_date():
    """Requirement 19.1: a date is a real date cell, sortable chronologically, not text."""
    workbook = openpyxl.load_workbook(io.BytesIO(to_xlsx(_sample_document())))
    sheet = workbook.active

    date_cells = [
        c
        for row in sheet.iter_rows()
        for c in row
        if isinstance(c.value, datetime) and c.value.year == 2025 and c.value.month == 8
    ]
    assert date_cells, "the date value did not land as a real date cell"
    assert date_cells[0].value.day == 4


def test_excel_integer_and_hours_cells_are_numeric():
    """Requirement 19.1: integer and hours columns are numeric so they sum."""
    workbook = openpyxl.load_workbook(io.BytesIO(to_xlsx(_sample_document())))
    sheet = workbook.active

    values = [c.value for row in sheet.iter_rows() for c in row if c.value is not None]
    assert 12720 in values, "the integer minutes value must be a number"
    assert any(isinstance(v, int | float) and abs(v - 212.0) < 1e-9 for v in values)


def test_excel_hebrew_sheet_is_right_to_left():
    """Requirement 19.3: a Hebrew workbook lays its columns out right-to-left."""
    workbook = openpyxl.load_workbook(io.BytesIO(to_xlsx(_sample_document("he"))))
    assert workbook.active.sheet_view.rightToLeft is True

    english = openpyxl.load_workbook(io.BytesIO(to_xlsx(_sample_document("en"))))
    assert not english.active.sheet_view.rightToLeft


def test_excel_carries_the_stamp():
    """Requirement 19.4: the workbook is stamped with period, generation time and user."""
    workbook = openpyxl.load_workbook(io.BytesIO(to_xlsx(_sample_document())))
    text = "\n".join(
        str(c.value) for row in workbook.active.iter_rows() for c in row if c.value is not None
    )
    assert "08/2025" in text
    assert "admin" in text
    assert "01/09/2025 10:30" in text


# ===================================================================== PDF: RTL Hebrew (19.2, 19.3)


def test_pdf_html_is_right_to_left_for_hebrew():
    """Requirement 19.3: the Hebrew PDF's HTML declares right-to-left and the Hebrew language."""
    markup = to_pdf_html(_sample_document("he"))
    assert 'dir="rtl"' in markup
    assert 'lang="he"' in markup
    # The Hebrew text is present in the document, so WeasyPrint lays out real Hebrew, not a placeholder.
    assert _HEB_NAME in markup
    assert _HEB_EMPLOYEE in markup


def test_pdf_html_is_left_to_right_for_english():
    """The English PDF is left-to-right, from the same document built with `language='en'`."""
    markup = to_pdf_html(_sample_document("en"))
    assert 'dir="ltr"' in markup
    assert 'lang="en"' in markup
    assert "Employee" in markup


def test_pdf_html_carries_the_stamp():
    """Requirement 19.4: the PDF is stamped with period, generation time and user."""
    markup = to_pdf_html(_sample_document())
    assert "08/2025" in markup
    assert "admin" in markup
    assert "01/09/2025 10:30" in markup


def test_pdf_html_escapes_markup_in_values():
    """A value containing angle brackets is escaped, so it cannot inject an element."""
    doc = ExportDocument(
        title="T",
        title_en="T",
        stamp=ExportStamp(period_label="08/2025", generated_at=datetime(2025, 9, 1), generated_by="admin"),
        tables=[
            ExportTable(
                columns=[ExportColumn("A", "A")],
                rows=[[cell("<script>x</script>")]],
            )
        ],
    )
    markup = to_pdf_html(doc)
    assert "<script>x</script>" not in markup
    assert "&lt;script&gt;" in markup


def test_pdf_bytes_render_when_the_native_stack_is_present():
    """Requirement 19.2: the PDF renders to bytes where WeasyPrint's libraries are available.

    On a host without Pango/cairo (a bare developer machine) `to_pdf` raises `PdfRenderUnavailable`;
    the test skips there rather than failing, because the RTL correctness it would assert lives in the
    HTML, which `test_pdf_html_*` cover without a native library. Where the stack is present — the
    container the app deploys to — the bytes are a real PDF.
    """
    try:
        pdf = to_pdf(_sample_document("he"))
    except PdfRenderUnavailable:
        pytest.skip("WeasyPrint native libraries not present on this host")
    assert pdf[:5] == b"%PDF-"
    assert len(pdf) > 500
